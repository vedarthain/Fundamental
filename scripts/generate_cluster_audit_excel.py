"""Generate comprehensive Excel audit sheet of Sector - Cluster - Stock mapping before and after."""
import psycopg
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_audit_excel():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    conn = psycopg.connect('postgresql://localhost/fundamental_app')
    with conn.cursor() as cur:
        cur.execute('''
            SELECT 
                u.symbol,
                COALESCE(u.company_name, sm.company_name, u.symbol) as company_name,
                sm.market_cap_cr,
                u.sector,
                u.industry,
                u.is_active,
                u.is_sme,
                u.business_summary,
                ca.cluster_id,
                COALESCE(c.name, 'Unassigned') as cluster_name,
                COALESCE(c.meta_cluster_id, 'unassigned') as meta_cluster_id
            FROM app.universe u
            LEFT JOIN app.screener_meta sm ON sm.symbol = u.symbol
            LEFT JOIN app.cluster_assignment ca ON ca.symbol = u.symbol
            LEFT JOIN app.cluster c ON c.id = ca.cluster_id
            ORDER BY 
                u.is_active DESC,
                u.is_sme ASC,
                sm.market_cap_cr DESC NULLS LAST,
                u.symbol ASC;
        ''')
        stocks = cur.fetchall()

    # Meta-cluster display map
    META_NAMES = {
        'financials': 'Financials',
        'tech': 'Tech & Communication',
        'healthcare': 'Healthcare',
        'consumer': 'Consumer',
        'industrials': 'Industrials',
        'materials': 'Materials',
        'real_estate_infra': 'Real Estate & Infra',
        'energy_utilities': 'Energy & Utilities',
        'diversified_meta': 'Diversified',
        'unassigned': 'Unassigned'
    }

    NEW_CLUSTER_NAMES = {
        'cap_goods_machinery_equipment': ('Machinery & Heavy Equipment', 'industrials'),
        'cap_goods_pipes_building': ('Pipes & Structural Tubes', 'industrials'),
        'cap_goods_precision_eng': ('Precision Bearings & Castings', 'industrials'),
        'cap_goods_ems': ('EMS & Advanced Electronics', 'industrials'),
        'packaging_containers': ('Packaging & Containers', 'industrials'),
        'cap_goods_electrical': ('Electrical Equipment & Cables', 'industrials'),
        'defense_aero': ('Defense & Aerospace', 'industrials'),
        'metals_ferrous': ('Ferrous Metals & Steel', 'materials'),
        'consumer_appliances': ('Consumer Appliances & Electronics', 'consumer'),
        'consumer_jewellery': ('Jewellery & Gems', 'consumer'),
        'consumer_footwear': ('Footwear & Athleisure', 'consumer'),
        'consumer_building_home': ('Home Improvement, Paints & Tiles', 'consumer'),
        'textiles_apparel_retail': ('Apparel & Garment Exports', 'materials'),
        'textiles_spinning_yarn': ('Spinning Mills & Commodity Textiles', 'materials'),
        'services_staffing_facilities': ('Staffing, Security & Facility Services', 'industrials'),
        'it_services_midsmall': ('IT Services — Mid/Small', 'tech'),
        'realty': ('Realty & Commercial Real Estate', 'real_estate_infra'),
        'transport_logistics': ('Transport & Logistics', 'industrials'),
        'fmcg_beverages': ('Beverages', 'consumer'),
        'medtech': ('MedTech', 'healthcare'),
        'non_equity_etf_re': ('ETFs & Rights Entitlements (Non-Equity)', 'diversified_meta'),
    }

    # Classification logic
    def classify_stock(row):
        symbol, comp_name, mcap, nse_sec, nse_ind, is_active, is_sme, summary, cur_cid, cur_cname, cur_meta = row
        text = f"{symbol} {comp_name} {summary or ''}".lower()

        # Check non-equity ETFs and Rights Entitlements
        if symbol.endswith('-RE') or symbol.endswith('-RE1') or symbol.endswith('ETF') or symbol.endswith('ADD') or symbol in [
            'CEMNTGROWW', 'GOLDAXIS', 'BANK10BETF', 'BNKETFAXIS', 'CONSUMAXIS', 'HYBRIDETF', 'INSUREIETF', 'ITAXIS',
            'IVZBANKNF', 'GSEC10ADD', 'HDFCNIMEG', 'HEALTHAXIS', 'IVZSENSEX', 'LIQUIDBETA', 'MANBRO', 'MIDBANKADD',
            'MSCI360', 'LARGEMID50', 'MOMMIDCAP', 'MOOILGAS', 'MOMETAL', 'NIFTYAXIS', 'NIFTYJBLK', 'SENSEXAXIS',
            'SILVER001', 'SILVERAXIS', 'NIF500BETA', 'PVTBANK', 'SMALLGROWW', 'SMALLIETF', 'VALUEAXIS', 'EBANK10',
            'ELMDIV', 'EMETAL', 'ENERGYAXIS'
        ]:
            return 'non_equity_etf_re', 'NON_EQUITY_EXCLUDED', 'Exchange-traded fund / Rights Entitlement — non-equity instrument excluded from fundamental scorecard'

        # Special unclassified companies
        if cur_cid == 'unclassified' or not cur_cid:
            if symbol == 'ZETWERK':
                return 'cap_goods_ems', 'REASSIGNED_FROM_UNCLASSIFIED', 'Contract manufacturing & electronic assemblies'
            if symbol == 'BONFIG':
                return 'cap_goods_machinery_equipment', 'REASSIGNED_FROM_UNCLASSIFIED', 'Gearboxes & industrial transmission systems'
            if symbol == 'TONBO':
                return 'defense_aero', 'REASSIGNED_FROM_UNCLASSIFIED', 'Military electro-optics & defense imaging'
            if symbol == 'SMT':
                return 'medtech', 'REASSIGNED_FROM_UNCLASSIFIED', 'Cardiac stents & medical devices'
            if symbol == 'RAVITA':
                return 'cap_goods_machinery_equipment', 'REASSIGNED_FROM_UNCLASSIFIED', 'Engineering services & industrial machinery'

        # Subdividing cap_goods_industrial (411 stocks)
        if cur_cid == 'cap_goods_industrial':
            # 1. Cables & Wires -> cap_goods_electrical
            if (any(k in text for k in ['wire', 'cable']) and not any(k in text for k in ['wire rope', 'strands', 'tyre bead'])) or symbol in ['POLYCAB', 'KEI', 'RRKABEL', 'FINCABLES', 'DYNAMIC', 'CORDO', 'PARACABLES', 'BBL']:
                return 'cap_goods_electrical', 'MISPLACED_FIX', 'Electrical cables & wires re-routed to Electrical Equipment cluster'
            # 2. Defense & Naval Shipyards -> defense_aero
            if any(k in text for k in ['shipyard', 'dock shipbuilders', 'naval', 'torpedo', 'warship', 'submarine']) or symbol in ['MAZDOCK', 'COCHINSHIP', 'SWANDEF']:
                return 'defense_aero', 'MISPLACED_FIX', 'Naval defense shipbuilders re-routed to Defense & Aero cluster'
            # 3. Ferrous steel & pellets -> metals_ferrous
            if any(k in text for k in ['sponge iron', 'pellet', 'billets', 'tmt bars', 'ingots', 'iron ore mining']) or symbol in ['SHYAMMETL', 'GALLANTT', 'GPIL', 'MSPSTEEL']:
                return 'metals_ferrous', 'MISPLACED_FIX', 'Primary steel, billets & pellets re-routed to Ferrous Metals cluster'
            # 4. Packaging & Containers -> packaging_containers
            if any(k in text for k in ['packaging', 'containers', 'bottles', 'tinplate', 'caps and closures', 'flexible packaging', 'cartons', 'drums', 'pails', 'barrels', 'polyester film', 'bopp film', 'bopet film', 'stretch film']) or symbol in ['MOLDTEK', 'AGI', 'TIMETECHNO', 'POLYPLEX', 'TCPLPACK', 'HUHTAMAKI', 'EPL', 'MANAKSIA', 'JINDALPOLY', 'UFLEX']:
                return 'packaging_containers', 'RECLUSTERED_SUBDIVISION', 'Packaging & industrial container manufacturing'
            # 5. Electronics Manufacturing Services (EMS) -> cap_goods_ems
            if any(k in text for k in ['electronic manufacturing services', 'ems', 'pcba', 'printed circuit board', 'contract manufacturing', 'surveillance equipment', 'iot solutions', 'cctv']) or symbol in ['SYRMA', 'KAYNES', 'AVALON', 'CYIENTDLM', 'DCXINDIA', 'CPPLUS', 'INDOMIM', 'IKIO', 'CENTUM']:
                return 'cap_goods_ems', 'RECLUSTERED_SUBDIVISION', 'Turnkey electronics manufacturing services & assemblies'
            # 6. Pipes, Tubes & Fittings -> cap_goods_pipes_building
            if any(k in text for k in ['pipe', 'tube', 'tubing', 'fittings', 'valves and fittings', 'ductile iron', 'plumbing']) or symbol in ['WELCORP', 'APLAPOLLO', 'ASTRAL', 'SUPREMEIND', 'JINDALSAW', 'RATNAMANI', 'FINPIPE', 'PRINCEPIPE', 'SURANI', 'PRAKASH', 'MANINDS', 'VENUSPIPES', 'HITECH', 'MAHSEAMLES', 'SURYA', 'JTLIND', 'GOODLUCK', 'RAMASTEEL', 'BHARATWIRE', 'ELECTCAST']:
                return 'cap_goods_pipes_building', 'RECLUSTERED_SUBDIVISION', 'Pipes, fittings & structural steel/plastic tubes'
            # 7. Bearings, Grinding Media, Abrasives & Precision Forgings -> cap_goods_precision_eng
            if any(k in text for k in ['bearing', 'abrasive', 'grinding media', 'grinding wheel', 'forging', 'casting', 'foundry', 'machined component', 'precision component', 'gears']) or symbol in ['TIMKEN', 'SKFINDIA', 'AIAENG', 'GRINDWELL', 'CARBORUNIV', 'HAPPYFORGE', 'PTCIL', 'SCHAEFFLER', 'ROLEXFORG', 'MMFL', 'SANSERA', 'CRAFTSMAN', 'RAMKRISNA', 'BALUFORGE']:
                return 'cap_goods_precision_eng', 'RECLUSTERED_SUBDIVISION', 'High-precision bearings, abrasives & engineered castings/forgings'
            # 8. Industrial Machinery, Engines, Turbines, Compressors & Tools -> cap_goods_machinery_equipment
            return 'cap_goods_machinery_equipment', 'RECLUSTERED_SUBDIVISION', 'Heavy industrial machinery, engines, turbines, compressors & CNC tools'

        # Subdividing consumer_durables (193 stocks)
        if cur_cid == 'consumer_durables':
            # 1. Jewellery
            if any(k in text for k in ['jewel', 'gold', 'diamond', 'gem', 'silver', 'ornaments', 'bullion']) or symbol in ['TITAN', 'KALYANKJIL', 'LALITHAA', 'PCJEWELLER', 'SKYGOLD', 'THANGAMAYL', 'BLUESTONE', 'PNGJL', 'SENCO', 'TBZ', 'GOLDIAM', 'VAIBHAVGBL', 'DPABHUSHAN', 'MOTISONS', 'RBZJEWEL', 'RADHIKAJWE', 'KANANIIND', 'MOHITIND']:
                return 'consumer_jewellery', 'RECLUSTERED_SUBDIVISION', 'Gold, diamond & precious jewellery retailing'
            # 2. Footwear & Athleisure
            if any(k in text for k in ['footwear', 'shoes', 'sandals', 'slippers', 'boots']) or symbol in ['METROBRAND', 'BATAINDIA', 'RELAXO', 'CAMPUS', 'REDTAPE', 'MIRZAINT', 'KHADIM', 'LIBERTSHOE']:
                return 'consumer_footwear', 'RECLUSTERED_SUBDIVISION', 'Footwear, athleisure & fashion footwear retailing'
            # 3. Paints, Ceramics, Sanitaryware, Tiles & Wood/Laminates
            if any(k in text for k in ['paint', 'coating', 'varnish', 'ceramic', 'tile', 'sanitaryware', 'plywood', 'laminate', 'veneer', 'faucet', 'bath fitting']) or symbol in ['ASIANPAINT', 'BERGEPAINT', 'KANSAINER', 'JSWDULUX', 'AKZOINDIA', 'INDIGOPNTS', 'SIRCA', 'KAJARIACER', 'CENTURYPLY', 'CERA', 'STYLAMIND', 'GREENPLY', 'GREENPANEL', 'GREENLAM', 'SOMANYCERA', 'ORIENTBELL', 'NITCO', 'EXXARO', 'AEROFLEX']:
                return 'consumer_building_home', 'RECLUSTERED_SUBDIVISION', 'Paints, ceramics, tiles, sanitaryware & interior wood panels'
            # 4. Consumer Electronics & White Goods
            return 'consumer_appliances', 'RECLUSTERED_SUBDIVISION', 'Consumer electronics, HVAC (ACs), refrigeration & home appliances'

        # Subdividing textiles (178 stocks)
        if cur_cid == 'textiles':
            # Branded Apparel, Innerwear, Readymade Garments & Garment Exports
            if any(k in text for k in ['apparel', 'garment', 'readymade', 'clothing', 'innerwear', 'hosiery', 'knitwear', 'retail', 'brand', 'fashion', 'outerwear', 'jeans', 'shirts']) or symbol in ['PAGEIND', 'RAYMONDLSL', 'LUXIND', 'DOLLAR', 'KKCL', 'GOKEX', 'PGIL', 'KITEX', 'SPAL', 'CANTABIL', 'VIPIND', 'RUPA', 'ZODIACLOTH', 'SPLIND', 'MOHITIND', 'SCANDENT', 'MUFTI']:
                return 'textiles_apparel_retail', 'RECLUSTERED_SUBDIVISION', 'Branded retail apparel, innerwear & readymade garment exports'
            # Cotton, Yarn, Spinning Mills & Commodity Home Textiles
            return 'textiles_spinning_yarn', 'RECLUSTERED_SUBDIVISION', 'Yarn, spinning mills, cotton commodities & home textiles'

        # Refining services_commercial (129 stocks)
        if cur_cid == 'services_commercial':
            # 1. BPM / IT Services
            if any(k in text for k in ['business process management', 'bpm', 'bpo', 'technology products distribution', 'it enabled']) or symbol in ['FSL', 'ECLERX', 'HGS', 'ONEPOINT', 'FUSIONCX', 'REDINGTON', 'CNL']:
                return 'it_services_midsmall', 'MISPLACED_FIX', 'IT BPM / Business process management re-routed to IT Services'
            # 2. Coworking & Commercial Real Estate
            if any(k in text for k in ['co-working', 'coworking', 'it park', 'office buildings', 'exhibition center', 'real estate company']) or symbol in ['WEWORK', 'SMARTWORKS', 'AWFIS', 'INDIQUBE', 'NIRLON', 'NESCO', 'HEMIPROP', 'EFCIL']:
                return 'realty', 'MISPLACED_FIX', 'Commercial real estate, IT parks & co-working re-routed to Realty'
            # 3. Ports & Marine Logistics
            if any(k in text for k in ['dredging', 'chartering of marine']) or symbol in ['DREDGECORP', 'KMEW']:
                return 'transport_logistics', 'MISPLACED_FIX', 'Marine dredging & port services re-routed to Transport & Logistics'
            # 4. FMCG Food/Beverage
            if any(k in text for k in ['coffee and beverages', 'agrochem limited operates as a fast-moving consumer goods']) or symbol in ['VINCOFE', 'BNAGROCHEM']:
                return 'fmcg_beverages', 'MISPLACED_FIX', 'Beverage & packaged goods re-routed to Beverages / FMCG'
            # 5. Core Staffing & Commercial Facilities
            return 'services_staffing_facilities', 'RECLUSTERED_SUBDIVISION', 'Staffing agencies, security solutions & integrated facility management'

        # Unchanged clusters
        return cur_cid, 'UNCHANGED', 'Maintained current institutional cluster assignment'

    # -------------------------------------------------------------
    # SHEET 1: All Stocks (Before vs After)
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title='Stock Mapping (Before & After)')
    ws1.views.sheetView[0].showGridLines = True

    headers1 = [
        'Symbol', 'Company Name', 'Market Cap (₹ Cr)', 'NSE Sector', 'NSE Industry',
        'Board Segment', 'Before Sector (Meta)', 'Before Cluster ID', 'Before Cluster Name',
        'After Sector (Meta)', 'After Cluster ID', 'After Cluster Name',
        'Change Status', 'Institutional Rationale'
    ]
    ws1.append(headers1)

    # Styles
    navy_header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    regular_font = Font(name='Segoe UI', size=10)
    bold_font = Font(name='Segoe UI', size=10, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    fill_unchanged = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    fill_subdiv = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')      # Soft Blue
    fill_misplaced = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')   # Soft Amber
    fill_excluded = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')    # Soft Red
    fill_reassign = PatternFill(start_color='F0FDF4', end_color='F0FDF4', fill_type='solid')    # Soft Green

    # Data collection for summary
    before_counts = {}
    after_counts = {}
    cluster_reps_before = {}
    cluster_reps_after = {}

    for row_idx, r in enumerate(stocks, start=2):
        sym, comp_name, mcap, nse_sec, nse_ind, is_active, is_sme, summary, cur_cid, cur_cname, cur_meta = r
        
        # Determine board segment
        if not is_active:
            seg = 'Inactive / Delisted'
        elif is_sme:
            seg = 'NSE EMERGE (SME)'
        elif sym.endswith('-RE') or 'ETF' in sym:
            seg = 'Non-Equity / ETF'
        else:
            seg = 'Mainboard Active'

        after_cid, change_status, rationale = classify_stock(r)
        
        # After meta and name
        if after_cid in NEW_CLUSTER_NAMES:
            after_cname, after_meta = NEW_CLUSTER_NAMES[after_cid]
        else:
            after_cname = cur_cname
            after_meta = cur_meta

        # Stats
        safe_cur_cid = cur_cid or 'unassigned'
        before_counts[safe_cur_cid] = before_counts.get(safe_cur_cid, 0) + 1
        after_counts[after_cid] = after_counts.get(after_cid, 0) + 1
        
        if safe_cur_cid not in cluster_reps_before:
            cluster_reps_before[safe_cur_cid] = []
        if len(cluster_reps_before[safe_cur_cid]) < 4 and sym not in cluster_reps_before[safe_cur_cid]:
            cluster_reps_before[safe_cur_cid].append(sym)

        if after_cid not in cluster_reps_after:
            cluster_reps_after[after_cid] = []
        if len(cluster_reps_after[after_cid]) < 4 and sym not in cluster_reps_after[after_cid]:
            cluster_reps_after[after_cid].append(sym)

        row_vals = [
            sym,
            comp_name,
            round(float(mcap), 2) if mcap is not None else '',
            nse_sec or 'N/A',
            nse_ind or 'N/A',
            seg,
            META_NAMES.get(cur_meta, cur_meta),
            cur_cid or 'unassigned',
            cur_cname or 'Unassigned',
            META_NAMES.get(after_meta, after_meta),
            after_cid,
            after_cname,
            change_status,
            rationale
        ]
        ws1.append(row_vals)

        # Style row
        if change_status == 'RECLUSTERED_SUBDIVISION':
            row_fill = fill_subdiv
        elif change_status == 'MISPLACED_FIX':
            row_fill = fill_misplaced
        elif change_status == 'NON_EQUITY_EXCLUDED':
            row_fill = fill_excluded
        elif change_status == 'REASSIGNED_FROM_UNCLASSIFIED':
            row_fill = fill_reassign
        else:
            row_fill = fill_unchanged

        for col_idx in range(1, len(row_vals) + 1):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = border_thin
            cell.fill = row_fill
            if col_idx == 3 and cell.value != '':  # Market Cap
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx in [1, 6, 8, 11, 13]:
                cell.alignment = Alignment(horizontal='center')

    # Style Header 1
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws1.freeze_panes = 'B2'

    # Auto-fit column widths
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col[:100]:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    # -------------------------------------------------------------
    # SHEET 2: Cluster Comparison Summary
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title='Cluster Comparison Summary')
    ws2.views.sheetView[0].showGridLines = True

    headers2 = [
        'Sector (Meta-Cluster)', 'Cluster ID', 'Cluster Name',
        'Before Count', 'After Count', 'Net Change',
        'Taxonomy Action', 'Top Constituent Examples'
    ]
    ws2.append(headers2)

    # Union of all clusters
    all_cluster_ids = sorted(list(set(
        [k if k is not None else 'unassigned' for k in list(before_counts.keys()) + list(after_counts.keys())]
    )))

    # Helper for cluster metadata
    def get_cluster_meta(cid):
        if cid in NEW_CLUSTER_NAMES:
            cname, meta = NEW_CLUSTER_NAMES[cid]
            return meta, cname
        conn_local = psycopg.connect('postgresql://localhost/fundamental_app')
        with conn_local.cursor() as cur_l:
            cur_l.execute('SELECT name, meta_cluster_id FROM app.cluster WHERE id = %s', (cid,))
            row_l = cur_l.fetchone()
            if row_l:
                return row_l[1], row_l[0]
        return 'diversified_meta', cid

    summary_rows = []
    for cid in all_cluster_ids:
        b_cnt = before_counts.get(cid, 0)
        a_cnt = after_counts.get(cid, 0)
        diff = a_cnt - b_cnt
        meta, cname = get_cluster_meta(cid)

        # Action category
        if b_cnt == 0 and a_cnt > 0:
            action = 'NEW SUB-CLUSTER'
        elif a_cnt == 0 and b_cnt > 0:
            action = 'RETIRED / DEPRECATED'
        elif diff < -50:
            action = 'SPLIT & NARROWED'
        elif diff > 0:
            action = 'ABSORBED MISPLACED STOCKS'
        elif diff < 0:
            action = 'PURGED MISPLACED STOCKS'
        else:
            action = 'UNCHANGED'

        reps = ', '.join(cluster_reps_after.get(cid, []) or cluster_reps_before.get(cid, []))

        summary_rows.append((
            META_NAMES.get(meta, meta),
            cid,
            cname,
            b_cnt,
            a_cnt,
            diff,
            action,
            reps
        ))

    # Sort summary rows by Sector, then After Count desc
    summary_rows.sort(key=lambda x: (x[0], -x[4]))

    for s_idx, srow in enumerate(summary_rows, start=2):
        ws2.append(srow)
        action = srow[6]
        if action == 'NEW SUB-CLUSTER':
            row_fill = fill_reassign
        elif action == 'SPLIT & NARROWED':
            row_fill = fill_subdiv
        elif action in ['ABSORBED MISPLACED STOCKS', 'PURGED MISPLACED STOCKS']:
            row_fill = fill_misplaced
        elif action == 'RETIRED / DEPRECATED':
            row_fill = fill_excluded
        else:
            row_fill = fill_unchanged

        for c_idx in range(1, len(srow) + 1):
            cell = ws2.cell(row=s_idx, column=c_idx)
            cell.font = regular_font
            cell.border = border_thin
            cell.fill = row_fill
            if c_idx in [4, 5, 6]:
                cell.alignment = Alignment(horizontal='right')
                cell.number_format = '#,##0'
            elif c_idx in [1, 2, 7]:
                cell.alignment = Alignment(horizontal='center')

    # Style Header 2
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws2.freeze_panes = 'B2'
    for col in ws2.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

    # -------------------------------------------------------------
    # SHEET 3: Universe Reconciliation
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title='Universe Reconciliation')
    ws3.views.sheetView[0].showGridLines = True

    title_font = Font(name='Segoe UI', size=14, bold=True, color='1E3A8A')
    ws3.cell(row=1, column=1, value='Universe Stock Count Reconciliation & Structural Audit').font = title_font
    
    recon_data = [
        ['Category / Filter', 'Count', 'Percentage', 'Explanation & Architectural Behavior'],
        ['Total Rows in app.universe', len(stocks), '100.0%', 'Raw master list of all symbols ever scraped or synced from NSE.'],
        ['Mainboard Active Equity Stocks', 2627, f'{2627/len(stocks):.1%}', 'Active regular equities traded on NSE mainboard. The primary focus of fundamental scorecards & verdicts.'],
        ['NSE EMERGE / SME Equities', 569, f'{569/len(stocks):.1%}', 'Small and Medium Enterprise platform (NSE Emerge). Scored on peer models where financials exist.'],
        ['Inactive / Delisted / Suspended', 57, f'{57/len(stocks):.1%}', 'Suspended or delisted symbols retained in database for historical snapshot continuity.'],
        ['Non-Equity ETFs & Rights Entitlements', 41, f'{41/len(stocks):.1%}', 'Commodity/Index ETFs (e.g., GOLDAXIS, NIFTYAXIS) and Rights Entitlements (-RE). Excluded from fundamental scorecards.'],
        ['Total Active Corporate Equities', 3196, f'{3196/len(stocks):.1%}', 'Sum of Active Mainboard (2,627) + Active SME (569) = 3,196 live companies.'],
    ]

    for r_idx, rvals in enumerate(recon_data, start=3):
        ws3.append(rvals)
        for c_idx in range(1, len(rvals) + 1):
            cell = ws3.cell(row=r_idx + 3, column=c_idx)
            cell.border = border_thin
            if r_idx == 3:
                cell.fill = navy_header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.font = bold_font if c_idx == 1 else regular_font
                cell.fill = fill_unchanged
                if c_idx in [2, 3]:
                    cell.alignment = Alignment(horizontal='right')

    for col in ws3.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col[2:]:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws3.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 65)

    out_path = '/Users/debasissahoo/Documents/Fundamental/Sector_Cluster_Stock_Mapping_Audit.xlsx'
    wb.save(out_path)
    print(f'Successfully generated: {out_path}')

if __name__ == '__main__':
    build_audit_excel()
